"""Template round trip: build, fill like a user, read back with validation."""
import io, os, sys, unittest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from openpyxl import load_workbook
from api.excel_io import template_xlsx, parse_xlsx

D = {'month': '2026-09', 'projects': [{'id': 'p1', 'name': 'Tower'}, {'id': 'p2', 'name': 'Plant'}],
     'monthly': [['empHours', 'Employee man-hours'], ['conHours', 'Contractor man-hours'], ['nhRecycled', 'Recycled (t)']],
     'env': [['electricity', 'Electricity (kWh)'], ['dieselL', 'Diesel (L)']],
     'existing': {'monthly': {'p1': {'empHours': 1000}}}}

class T(unittest.TestCase):
    def test_round_trip(self):
        wb = load_workbook(io.BytesIO(template_xlsx(D)))
        self.assertEqual(wb.sheetnames, ['Read me', 'Monthly data', 'Environmental'])
        m = wb['Monthly data']
        self.assertEqual(m['D3'].value, 1000)          # existing figure prefilled
        m['E3'] = 2500; m['D4'] = '3,000'; m['F4'] = -2   # p2 has a negative value
        e = wb['Environmental']; e['D3'] = 120000; e['E4'] = 'abc'
        m.append(['ghost', 'Unknown', '2026-09', 5])
        b = io.BytesIO(); wb.save(b)
        r = parse_xlsx(b.getvalue(), D)
        rows = {(x['kind'], x['project']): x['values'] for x in r['rows']}
        self.assertEqual(rows[('monthly', 'p1')], {'empHours': 1000, 'conHours': 2500})
        self.assertEqual(rows[('env', 'p1')], {'electricity': 120000})
        self.assertNotIn(('monthly', 'p2'), rows)        # rejected: negative
        self.assertNotIn(('env', 'p2'), rows)            # rejected: not a number
        self.assertTrue(any('negative' in x for x in r['errors']))
        self.assertTrue(any("'abc'" in x for x in r['errors']))
        self.assertTrue(any("ghost" in x for x in r['errors']))
    def test_bad_file(self):
        with self.assertRaises(ValueError):
            parse_xlsx(b'not excel', D)

if __name__ == '__main__':
    unittest.main()
