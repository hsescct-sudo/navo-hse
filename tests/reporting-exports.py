"""Exercise the real export builders in memory using the captured dashboard fixture."""
import io,json,sys,zipfile
from pathlib import Path
root=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(root),str(root.parent/'report-deps')]
from api.report import make
from pptx import Presentation
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
payload=json.loads((root/'tests/screens/captured-report-payload.json').read_text(encoding='utf-8'))
payload.update(kind='exec',brief=True,dashboardId='es',dashboardTitle='Environmental & ESG',reportSections=[{'title':'Environmental & ESG metrics','headers':['Metric','Recorded total','Unit','Records'],'rows':[[f'Metric {i}',i,'kWh',1] for i in range(14)]}])
for fmt in ['pptx','docx','xlsx','pdf']:
 data,mime=make(fmt,payload)
 assert len(data)>500
 if fmt=='pptx':
  prs=Presentation(io.BytesIO(data));assert len(prs.slides)==3
  assert any(shape.shape_type==13 for shape in prs.slides[1].shapes),'slide two must contain the dashboard'
  assert 'Metric 13' in '\n'.join(cell.text for shape in prs.slides[2].shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
 elif fmt=='docx':
  doc=Document(io.BytesIO(data));assert len(doc.inline_shapes)==1;assert len(doc.tables)==1;assert len(doc.tables[0].rows)==15
 elif fmt=='xlsx':
  book=load_workbook(io.BytesIO(data));assert 'Dashboard' in book.sheetnames;assert book['Environmental & ESG metrics'].max_row==15
 else:assert len(PdfReader(io.BytesIO(data)).pages)==1
case={'kind':'case','brand':'NAVO','title':'Wheel evidence validation','reference':7654321,'project':'Validation site','fields':[{'label':'Electrical / evidence','value':'Photo E1: isolation switch damaged'},{'label':'What source of energy was present?','value':'Electrical energy'},{'label':'How will we verify that the improvements are making a difference in practice?','value':'Supervisor witnesses field test'}],'actions':[{'title':'Restore isolation switch','owner':'Test supervisor','due':'2026-10-30','priority':'High'}]}
b,_=make('pdf',case);text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(b)).pages);normalized=' '.join(text.split());assert 'Photo E1' in normalized and 'Supervisor witnesses field test' in normalized,repr(text)
print('PASS real Word/Excel/PowerPoint/PDF builders: three slides, dashboard on slide two, all 14 ESG metrics, one-page dashboard PDF, investigation evidence in case PDF')
