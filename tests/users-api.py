"""Unit tests for api/users.py (no network)."""
import base64, io, os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from openpyxl import load_workbook
from api import users as U

b64 = U.template_xlsx(["System administrator", "HSE engineer / officer"], None, ["Power"], ["P-1", "P-2"])
wb = load_workbook(io.BytesIO(base64.b64decode(b64)))
assert wb.sheetnames[:2] == ["Accounts", "Lists"]
ws = wb["Accounts"]
ws.append(["Sara Omar", " Sara@X.com ", "HSE engineer / officer", "Project", "", "P-1; P-2", "Secret123"])
ws.append([None] * 7)
buf = io.BytesIO(); wb.save(buf)
rows = U.parse_accounts(buf.getvalue())
assert len(rows) == 2, rows
assert rows[1] == {"name": "Sara Omar", "email": "sara@x.com", "position": "HSE engineer / officer", "level": "Project", "scope": "", "projects": ["P-1", "P-2"], "password": "Secret123"}, rows[1]
try:
    U.parse_accounts(b"not excel"); raise SystemExit("bad file accepted")
except U.ApiError as e:
    assert e.status == 400
for _ in range(50):
    p = U.gen_password()
    assert len(p) == 12 and any(c.isdigit() for c in p) and any(c.isupper() for c in p)
try:
    U.check_password("short"); raise SystemExit("short accepted")
except U.ApiError:
    pass
# self-protection
U._http = lambda *a, **k: []
for fn, args in ((U.delete_user, ("me", "me")), (U.update_user, ({"id": "me", "disabled": True}, "me"))):
    try:
        fn(*args); raise SystemExit("self action allowed")
    except U.ApiError:
        pass
os.environ.pop("NAVO_SUPABASE_SERVICE_KEY", None)
print("users-api ok")
