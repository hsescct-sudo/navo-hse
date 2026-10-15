"""Presentation-ready reports. All numbers come from the supplied accessible dashboard payload."""
import io, math, textwrap
INK='#163747'; TEAL='#167563'; BLUE='#2765bb'; MUTED='#607986'; LIGHT='#edf3f5'; RED='#c24450'
def _plot():
 import matplotlib
 matplotlib.use('Agg')
 import matplotlib.pyplot as plt
 plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False,'axes.edgecolor':'#dce6eb','text.color':INK,'axes.labelcolor':MUTED,'xtick.color':MUTED,'ytick.color':MUTED,'savefig.facecolor':'#f5f8fa','pdf.fonttype':42})
 return plt

def _base(d,title,landscape=True):
 plt=_plot();fig=plt.figure(figsize=(14.4,8.1) if landscape else (8.27,11.69),facecolor='#f5f8fa')
 fig.text(.035,.955,str(d.get('brand','NAVO'))+' / LEADERSHIP INTELLIGENCE',fontsize=10,color=TEAL,weight='bold')
 fig.text(.035,.902,title,fontsize=24,weight='bold')
 fig.text(.035,.866,str(d.get('scope',d.get('project','')))+'  |  '+str(d.get('period',d.get('date',''))),fontsize=10,color=MUTED)
 flag='ILLUSTRATIVE DEMONSTRATION / NOT VERIFIED COMPANY PERFORMANCE' if d.get('illustrative') else 'Source: accessible field records / review before external issue'
 fig.text(.965,.958,flag,fontsize=7,color='#976b13' if d.get('illustrative') else MUTED,ha='right')
 fig.text(.035,.027,'Generated '+str(d.get('generated',''))+'  /  '+str(d.get('sourceUrl','https://navo-hse.vercel.app/')),fontsize=8,color=MUTED,url=d.get('sourceUrl','https://navo-hse.vercel.app/'))
 return fig

def _table(fig,rect,headers,rows,widths=None,size=9):
 ax=fig.add_axes(rect);ax.axis('off');t=ax.table(cellText=rows or [['—']*len(headers)],colLabels=headers,colWidths=widths,loc='upper center',cellLoc='left');t.auto_set_font_size(False);t.set_fontsize(size);t.scale(1,1.65)
 for (i,j),c in t.get_celld().items():
  c.set_edgecolor('#e1eaef');c.set_linewidth(.5);c.PAD=.09;c.set_facecolor(INK if i==0 else ('white' if i%2 else LIGHT));c.get_text().set_color('white' if i==0 else INK)
  if i==0:c.get_text().set_weight('bold')
 return t

def _num(x):return '—' if x is None else f'{x:,.0f}' if isinstance(x,(int,float)) else str(x)
def _rate(x):return 'Unavailable' if x is None else f'{x:.2f}'
def executive_pdf(d):
 plt=_plot();c=d.get('cur',{});fig=_base(d,'Executive safety overview');
 metrics=[('Projects',c.get('projects')),('People',c.get('manpower')),('Exposure hours',f"{c.get('hours',0)/1e6:.1f}M"),('Recordable cases',c.get('trc')),('Overdue actions',c.get('over'))]
 for i,(label,value) in enumerate(metrics):
  x=.035+i*.19;fig.text(x,.806,label,fontsize=9,color=MUTED);fig.text(x,.749,_num(value),fontsize=28,weight='bold',color=RED if i==4 else INK)
 fig.text(.035,.686,'PLATFORM PERFORMANCE',fontsize=10,weight='bold',color=TEAL)
 rates=d.get('platformRates') or [dict(name=d.get('scope','Corporation'),hours=c.get('hours'),**{k:c.get(k) if c.get('hours') else None for k in ['ltifr','ltisr','trcf']},targets=d.get('targets',{}))]
 _table(fig,[.035,.455,.45,.215],['Platform','LTIFR','LTISR','TRCF'],[[r['name'].replace('Non-Construction','Operations')]+[_rate(r.get(k))+' / '+_rate(r.get('targets',{}).get(k)) for k in ['ltifr','ltisr','trcf']] for r in rates],[.31,.23,.23,.23],9)
 fig.text(.035,.45,'Actual / target · rates per 1,000,000 exposure hours',fontsize=8,color=MUTED)
 fig.text(.53,.686,'FINDING RISK PROFILE',fontsize=10,weight='bold',color=TEAL)
 risk=[r for r in d.get('riskDistribution',[]) if r.get('count')];ax=fig.add_axes([.52,.425,.22,.245]);
 if risk:ax.pie([r['count'] for r in risk],colors=[r['color'] for r in risk],startangle=90,wedgeprops={'width':.3,'edgecolor':'white'},autopct='%1.0f%%',pctdistance=.82,textprops={'fontsize':8,'weight':'bold','color':'white'})
 else:ax.text(.5,.5,'No assessed data',ha='center',va='center');ax.axis('off')
 for i,r in enumerate(risk):fig.text(.755,.635-i*.035,r['label']+'  '+_num(r['count']),fontsize=9,color=r['color'])
 fig.text(.035,.388,'PRIORITY PROJECTS',fontsize=10,weight='bold',color=TEAL)
 projects=sorted(d.get('projects',[]),key=lambda x:(-(x.get('over') or 0),-(x.get('trcf') or 0)))[:6]
 _table(fig,[.035,.085,.45,.28],['Project','Division','TRCF','Overdue'],[[textwrap.shorten(str(x.get('name','')),width=31,placeholder='…'),textwrap.shorten(str(x.get('division','')),width=18,placeholder='…'),_rate(x.get('trcf')),_num(x.get('over'))] for x in projects],[.43,.27,.16,.14],8)
 fig.text(.53,.388,'FIVE-YEAR FREQUENCY TREND',fontsize=10,weight='bold',color=TEAL)
 rows=d.get('history',{}).get('rows',[]);ax=fig.add_axes([.56,.125,.39,.22]);
 for k,color in [('ltifr',BLUE),('trcf',TEAL)]:ax.plot([r['year'] for r in rows],[float('nan') if r.get(k) is None else r[k] for r in rows],marker='o',ms=4,label=k.upper(),color=color,linewidth=2)
 ax.set_xticks([r['year'] for r in rows]);ax.grid(axis='y',alpha=.15);ax.legend(frameon=False,fontsize=8);ax.tick_params(labelsize=8)
 note='2022–2025: simulated demonstration. Current year: stored YTD.' if d.get('history',{}).get('demoHistory') else 'Stored records only. Missing exposure produces unavailable rates.'
 fig.text(.53,.075,note,fontsize=7.5,color=MUTED)
 out=io.BytesIO();fig.savefig(out,format='pdf');plt.close(fig);return out.getvalue()

def history_pdf(d):
 plt=_plot();fig=_base(d,'Five-year safety performance');rows=d.get('rows',[])[:5];ax=fig.add_axes([.08,.43,.86,.36]);
 for k,col in [('ltifr',BLUE),('ltisr','#d39a37'),('trcf',TEAL)]:ax.plot([r['year'] for r in rows],[float('nan') if r.get(k) is None else r[k] for r in rows],marker='o',linewidth=2,label=k.upper(),color=col)
 ax.set_xticks([r['year'] for r in rows]);ax.set_ylabel('Rate per 1,000,000 hours');ax.grid(axis='y',alpha=.15);ax.legend(frameon=False,ncols=3)
 _table(fig,[.035,.115,.93,.25],['Year','Exposure hours','LTI','Lost days','Recordables','LTIFR','LTISR','TRCF'],[[str(r['year'])+(' DEMO' if r.get('simulated') else ''),_num(r.get('hours')) if r.get('hours') else 'Not recorded',_num(r.get('lti')),_num(r.get('lost')),_num(r.get('recordable'))]+[_rate(r.get(k)) for k in ['ltifr','ltisr','trcf']] for r in rows],[.13,.18,.08,.11,.13,.12,.12,.13],9)
 note=d.get('historyNote') or 'Missing exposure produces unavailable rates. Validate reporting boundaries and source records before tender use.'
 fig.text(.035,.075,'\n'.join(textwrap.wrap(note,145)),fontsize=8,color='#976b13' if d.get('demoHistory') else MUTED)
 out=io.BytesIO();fig.savefig(out,format='pdf');plt.close(fig);return out.getvalue()

def case_pdf(d):
 plt=_plot();from matplotlib.backends.backend_pdf import PdfPages
 sections=[('CASE SUMMARY',[('Reference',d.get('reference')),('Report type',d.get('type')),('Status',d.get('status')),('Project',d.get('project')),('Location',d.get('location')),('Risk rating',d.get('risk'))]),('OBSERVATION / EVENT',[('',d.get('description'))]),('IMMEDIATE CONTROL',[('',d.get('immediateAction'))]),('INVESTIGATION',[('Conclusion',d.get('cause')),('Review','Recorded' if d.get('reviewed') else 'Pending / not applicable')])]
 if d.get('fields'):sections.append(('FORM DETAILS',[(str(x.get('label','')),x.get('value')) for x in d['fields']]))
 if d.get('actions'):sections.append(('ACTION PLAN',[(str(i+1)+'. '+str(a.get('title','')),'Owner: '+str(a.get('owner') or 'Unassigned')+' | Due: '+str(a.get('due') or 'Not set')+' | Priority: '+str(a.get('priority') or 'Not set')) for i,a in enumerate(d['actions'])]))
 out=io.BytesIO();pdf=PdfPages(out);page=0;fig=None;y=0
 def newpage():
  nonlocal page,fig,y
  if fig is not None:pdf.savefig(fig);plt.close(fig)
  page+=1;fig=_base(d,'Case #'+str(d.get('reference','')),False);fig.text(.035,.812,'\n'.join(textwrap.wrap(str(d.get('title','')),72)),fontsize=12,weight='bold',va='top');fig.text(.965,.027,str(page),fontsize=8,ha='right',color=MUTED);y=.75
 for head,fields in sections:
  if y<.18:newpage()
  if fig is None:newpage()
  fig.text(.035,y,head,fontsize=10,weight='bold',color=TEAL);y-=.027
  for label,value in fields:
   if value is None or str(value).strip()=='':continue
   lines=textwrap.wrap((label+': ' if label else '')+str(value),96) or ['']
   for line in lines:
    if y<.08:newpage();fig.text(.035,y,head+' / CONTINUED',fontsize=10,weight='bold',color=TEAL);y-=.027
    fig.text(.035,y,line,fontsize=9,va='top');y-=.019
   y-=.009
  y-=.016
 if fig is None:newpage()
 pdf.savefig(fig);plt.close(fig);pdf.close();return out.getvalue()

def executive_pptx(d):
 from pptx import Presentation
 from pptx.util import Inches, Pt
 from pptx.dml.color import RGBColor
 from pptx.chart.data import CategoryChartData
 from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
 prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9);s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F5F8FA')
 def text(x,y,w,h,value,size=15,bold=False,color='163747',url=None):
  box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=box.text_frame;tf.word_wrap=True;p=tf.paragraphs[0];p.text=str(value);p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
  if url:
   for r in p.runs:r.hyperlink.address=url
  return box
 text(.5,.22,14,.3,str(d.get('brand','NAVO'))+' / LEADERSHIP OVERVIEW',12,True,'167563');text(.5,.7,14,.6,'Executive safety overview',30,True);text(.5,1.35,14,.3,str(d.get('scope',''))+' | '+str(d.get('period','')),13,False,'607986')
 c=d.get('cur',{})
 for i,(label,v) in enumerate([('PROJECTS',c.get('projects')),('PEOPLE',c.get('manpower')),('EXPOSURE HOURS',str(round(c.get('hours',0)/1e6,1))+'M'),('RECORDABLE CASES',c.get('trc')),('OVERDUE ACTIONS',c.get('over'))]):
  text(.5+i*3.05,1.98,2.9,.3,label,11,False,'607986');text(.5+i*3.05,2.35,2.9,.6,_num(v),30,True,'C24450' if i==4 else '163747')
 def table(headers,rows,x,y,w,h):
  t=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y), Inches(w), Inches(h)).table
  for i,row in enumerate([headers]+rows):
   for j,v in enumerate(row):
    cell=t.cell(i,j);cell.text=str(v);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('163747' if i==0 else 'FFFFFF');p=cell.text_frame.paragraphs[0];p.font.name='Arial';p.font.size=Pt(11);p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '163747');p.font.bold=i==0
 text(.5,3.18,7,.3,'PLATFORM PERFORMANCE / ACTUAL vs TARGET',12,True,'167563')
 rates=d.get('platformRates',[]);table(['Platform','LTIFR','LTISR','TRCF'],[[r['name'].replace('Non-Construction','Operations')]+[_rate(r.get(k))+' / '+_rate(r.get('targets',{}).get(k)) for k in ['ltifr','ltisr','trcf']] for r in rates],.5,3.65,7.2,1.6)
 text(.5,5.42,7,.3,'Rates per 1,000,000 exposure hours',11,False,'607986');text(.5,6,7,.3,'ACTION & LEADERSHIP ASSURANCE',12,True,'167563');table(['Open actions','Closed','Leadership engagements'],[[_num(c.get('open')),_num(c.get('closed')),_num(sum(p.get('count',0) for p in d.get('leadership',[])))]],.5,6.45,7.2,.9)
 hist=d.get('history',{});rows=hist.get('rows',[]);data=CategoryChartData();data.categories=[str(r['year']) for r in rows]
 for k in ['ltifr','trcf']:data.add_series(k.upper(),[r.get(k) for r in rows])
 chart=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(8.25), Inches(3.55), Inches(7.2), Inches(3.95),data).chart;chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.font.size=Pt(11);chart.category_axis.tick_labels.font.size=Pt(11);chart.value_axis.tick_labels.font.size=Pt(10)
 text(8.25,3.18,7,.3,'FIVE-YEAR PERFORMANCE',12,True,'167563');text(8.25,7.6,7,.5,'2022–2025 simulated demo history / 2026 stored YTD' if hist.get('demoHistory') else 'Stored records only; missing exposure creates gaps',11,False,'607986');text(.5,8.3,12,.3,'ILLUSTRATIVE DATA / NOT VERIFIED COMPANY PERFORMANCE' if d.get('illustrative') else 'Source: accessible field records',10,False,'976B13');text(12.5,8.3,3,.3,'Open live dashboard ↗',12,True,'167563',d.get('sourceUrl','https://navo-hse.vercel.app/'));out=io.BytesIO();prs.save(out);return out.getvalue()
