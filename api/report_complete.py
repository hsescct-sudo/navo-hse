"""Complete report exports shared by Word, Excel and PowerPoint."""
import io,base64,textwrap
from api.report_layout import executive_pdf, executive_pptx as summary_pptx

def image_bytes(d):
 value=d.get('dashboardImage','')
 if not value:return None
 if not value.startswith('data:image/jpeg;base64,'):raise ValueError('Dashboard image must be JPEG')
 b=base64.b64decode(value.split(',',1)[1],validate=True)
 if len(b)>1_500_000:raise ValueError('Dashboard image is too large')
 from PIL import Image
 im=Image.open(io.BytesIO(b))
 if im.format!='JPEG' or im.width>5000 or im.height>8000:raise ValueError('Invalid dashboard dimensions')
 return b

def brief_pdf(d):
 from PIL import Image
 b=image_bytes(d)
 if not b:return executive_pdf(d)
 image=Image.open(io.BytesIO(b)).convert('RGB')
 out=io.BytesIO();image.save(out,format='PDF',resolution=110);return out.getvalue()

def full_sections(d):
 c=d.get('cur',{});out=[]
 out.append(('Reporting scope',['Item','Value'],[['Platform',d.get('scope')],['Period',d.get('period')],['Country',d.get('country')],['Data status','Illustrative demonstration' if d.get('illustrative') else 'Stored source records'],['Historical coverage',d.get('history',{}).get('historyNote') or 'Stored figures; missing exposure gives unavailable rates.']]))
 out.append(('Workforce and exposure',['Platform','Employees','Contractors','Employee hours','Contractor hours','Safe hours','Safe days'],[[r.get(k) for k in ['name','employees','contractors','employeeHours','contractorHours','safeHours','safeDays']] for r in d.get('workforce',[])]))
 out.append(('Geographical exposure',['Country','Employees','Contractors','Employee hours','Contractor hours','Safe hours','Safe days'],[[r.get(k) for k in ['country','employees','contractors','employeeHours','contractorHours','safeHours','safeDays']] for r in d.get('geography',[])]))
 out.append(('Frequency rates',['Platform','Exposure hours','LTIFR','LTISR','TRCF','LTIFR target','LTISR target','TRCF target'],[[r.get('name'),r.get('hours'),r.get('ltifr'),r.get('ltisr'),r.get('trcf')]+[r.get('targets',{}).get(k) for k in ['ltifr','ltisr','trcf']] for r in d.get('platformRates',[])]))
 out.append(('Safety outcomes and actions',['Indicator','Current','Previous'],[[label,c.get(key),d.get('prev',{}).get(key)] for label,key in [('Projects','projects'),('People','manpower'),('Exposure hours','hours'),('Fatalities','fat'),('Lost time injuries','lti'),('Lost workdays','lost'),('Recordable cases','trc'),('First aid','fac'),('Near misses','nm'),('Inspections','insp'),('Observations','obs'),('Training hours','trainH'),('Audits and findings','audits'),('Actions','acts'),('Open actions','open'),('Closed actions','closed'),('Overdue actions','over'),('Closure %','closePct')]]))
 out.append(('Leadership engagement',['Programme','Recorded engagements'],[[r.get('name'),r.get('count')] for r in d.get('leadership',[])]))
 out.append(('Monthly trends',['Month','Exposure hours','Recordable cases','Previous recordables','LTIFR'],[[n]+[d.get(k,[])[i] if i<len(d.get(k,[])) else None for k in ['trendHours','trendTrc','trendTrcPrev','trendLtifr']] for i,n in enumerate(d.get('months',[]))]))
 out.append(('Division performance',['Division','Exposure hours','LTIFR','TRCF','LTI','Recordables','Inspections','Closure %'],[[r.get(k) for k in ['name','hours','ltifr','trcf','lti','trc','insp','closure']] for r in d.get('divisions',[])]))
 out.append(('Project performance',['Project','Division','Platform','Country','Hours','TRCF','Overdue','Closure %','Health'],[[r.get(k) for k in ['name','division','platform','country','hours','trcf','over','closure','health']] for r in d.get('projects',[])]))
 out.append(('Finding risk distribution',['Risk rating','Findings'],[[r.get('label'),r.get('count')] for r in d.get('riskDistribution',[])]))
 out.append(('Five-year performance',['Year','Exposure hours','LTI','Lost days','Recordables','LTIFR','LTISR','TRCF','Coverage'],[[r.get(k) for k in ['year','hours','lti','lost','recordable','ltifr','ltisr','trcf','coverage']] for r in d.get('history',{}).get('rows',[])]))
 if d.get('includeRegisters',False):
  out.append(('Source record register',['Reference','Title','Type','Project','Status','Reported'],[[r.get(k) for k in ['reference','title','type','project','status','date']] for r in d.get('register',[])]))
  out.append(('Action register',['Reference','Action','Project','Owner','Due','Status','Priority'],[[r.get(k) for k in ['reference','title','project','owner','due','status','priority']] for r in d.get('actionRegister',[])]))
 return out

def sections(d):
 if not d.get('brief'):return full_sections(d)
 if d.get('reportSections'):
  return [(str(s.get('title','Selected dashboard'))[:160],s.get('headers',[])[:8],[r[:8] for r in s.get('rows',[])[:6]]) for s in d['reportSections'][:1]]
 keys={'summary':['Safety outcomes and actions'],'portfolio':['Project performance'],'evidence':['Finding risk distribution'],'exec':['Safety outcomes and actions'],'health':['Project performance'],'site':['Project performance'],'inj':['Safety outcomes and actions'],'aud':['Safety outcomes and actions'],'acc':['Safety outcomes and actions'],'kpi':['Frequency rates'],'es':['Safety outcomes and actions']}
 selected=keys.get(d.get('dashboardId'),['Safety outcomes and actions'])
 result=[]
 for title,headers,rows in full_sections(d):
  if title not in selected:continue
  if title=='Safety outcomes and actions':
   names={'Projects','People','Exposure hours','Recordable cases','Open actions','Overdue actions'}
   if d.get('dashboardId') in ['inj','acc']:names={'Fatalities','Lost time injuries','Lost workdays','Recordable cases','First aid','Near misses'}
   elif d.get('dashboardId')=='aud':names={'Audits and findings','Inspections','Open actions','Overdue actions','Closure %'}
   elif d.get('dashboardId')=='es':names={'Projects','People'}
   rows=[r for r in rows if r[0] in names]
  if title=='Project performance':
   headers=['Project','Country','Overdue','Closure %','Health'];rows=[[r[i] for i in [0,3,6,7,8]] for r in rows[:6]]
  result.append((d.get('dashboardTitle',title)+' / summary',headers,rows[:6]))
 return result

def value(v):
 if v is None:return '—'
 if isinstance(v,float):return f'{v:,.2f}'
 if isinstance(v,int):return f'{v:,}'
 return str(v)

def complete_docx(d):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.enum.section import WD_SECTION_START,WD_ORIENT
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 doc=Document();
 for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor.from_string('163747')
 style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style.font.color.rgb=RGBColor.from_string('163747')
 for sec in doc.sections:sec.top_margin=sec.bottom_margin=Inches(.65);sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
 doc.add_paragraph(str(d.get('brand','NAVO'))+' / HSSE REPORT','Subtitle');doc.add_heading(d.get('dashboardTitle','Leadership')+' / executive brief',0);doc.add_paragraph(d.get('scope','Corporation')+' | '+d.get('period',''));doc.add_paragraph('Generated '+str(d.get('generated','')))
 doc.add_heading('Brief contents',1)
 for title,_,_ in sections(d):doc.add_paragraph(title,style='List Bullet')
 doc.add_paragraph('ILLUSTRATIVE DEMONSTRATION / NOT VERIFIED COMPANY PERFORMANCE' if d.get('illustrative') else 'Review reporting boundaries and evidence before external issue.')
 doc.add_paragraph(d.get('sourceUrl',''))
 sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.left_margin=sec.right_margin=Inches(.5)
 doc.add_heading(d.get('dashboardTitle','Dashboard'),1);b=image_bytes(d)
 if b:
  from PIL import Image
  image=Image.open(io.BytesIO(b));ratio=min(10.5/image.width,5.7/image.height);doc.add_picture(io.BytesIO(b),width=Inches(image.width*ratio),height=Inches(image.height*ratio))
 else:doc.add_paragraph('Dashboard snapshot unavailable in this request. Detailed source tables follow. Export from the leadership overview to include the live dashboard.')
 previous_wide=None
 for title,headers,rows in sections(d):
  wide=True
  if wide!=previous_wide:
   sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT;sec.page_width=Inches(11.69 if wide else 8.27);sec.page_height=Inches(8.27 if wide else 11.69);sec.left_margin=sec.right_margin=Inches(.55);previous_wide=wide
  else:doc.add_paragraph('')
  doc.add_heading(title,1);caption=doc.add_paragraph(f'{len(rows)} records | '+str(d.get('period','')),style='Caption');caption.paragraph_format.keep_with_next=True
  t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.autofit=False
  weights=[3 if h in ['Title','Action','Value','Programme','Project'] else 2 if h in ['Division','Type','Owner','Coverage','Health'] else 1 for h in headers];available=sec.page_width-sec.left_margin-sec.right_margin
  for col,weight in zip(t.columns,weights):col.width=int(available*weight/sum(weights))
  for c,h in zip(t.rows[0].cells,headers):
   c.text=h;shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'163747');c._tc.get_or_add_tcPr().append(shade)
   for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor.from_string('FFFFFF');r.bold=True
  repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
  for row in rows:
   for index,(cell,v) in enumerate(zip(t.add_row().cells,row)):
    cell.text=(v.split('T',1)[0]+'\n'+v.split('T',1)[1][:5]) if headers[index]=='Reported' and isinstance(v,str) and 'T' in v else value(v)
  for index,row in enumerate(t.rows):
   for cell in row.cells:
    for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=(index < len(t.rows)-1) if len(rows)<=20 else index==0
  if not rows:doc.add_paragraph('No records in the selected scope.')
 for sec in doc.sections:
  p=sec.footer.paragraphs[0];p.text=str(d.get('brand','NAVO'))+' / '+d.get('scope','')+' / Page ';f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
 for style in doc.styles:
  for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
 for paragraph in doc.paragraphs:
  for border in paragraph._p.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
 out=io.BytesIO();doc.save(out);return out.getvalue()

def complete_xlsx(d):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.chart import LineChart,Reference
 from openpyxl.drawing.image import Image
 from openpyxl.workbook.properties import CalcProperties
 wb=Workbook();wb.remove(wb.active);summary=wb.create_sheet('Report guide');summary.append([d.get('brand','NAVO')+' / COMPLETE HSSE REPORT']);summary.append(['Platform',d.get('scope')]);summary.append(['Period',d.get('period')]);summary.append(['Country',d.get('country')]);summary.append(['Data','Illustrative demonstration' if d.get('illustrative') else 'Stored records']);summary.append(['Live dashboard',d.get('sourceUrl')]);summary['B6'].hyperlink=d.get('sourceUrl','');dash=wb.create_sheet('Dashboard');b=image_bytes(d)
 if b:
  image=Image(io.BytesIO(b));ratio=min(1450/image.width,850/image.height);image.width*=ratio;image.height*=ratio;dash.add_image(image,'A1')
 else:dash.append(['Open the leadership overview to include the dashboard snapshot.'])
 for title,headers,rows in sections(d):
  ws=wb.create_sheet(title.replace('/','-')[:31]);ws.append(headers)
  for row in rows:ws.append(["'"+v if isinstance(v,str) and v[:1] in '=+-@' else v for v in row])
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for j,h in enumerate(headers,1):
   from openpyxl.utils import get_column_letter
   ws.column_dimensions[get_column_letter(j)].width=48 if h in ['Title','Action','Programme','Value'] else 28 if h in ['Project','Division','Owner'] else 20
  for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='163747');cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True)
  for row in ws.iter_rows(min_row=2):
   for cell in row:cell.font=Font(name='Arial',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True)
  ws.sheet_view.showGridLines=False
  if title=='Five-year performance' and rows:
   for i in range(2,len(rows)+2):
    for col,num in [('F','C'),('G','D'),('H','E')]:ws[f'{col}{i}']=f'=IF(B{i}>0,{num}{i}*1000000/B{i},"")';ws[f'{col}{i}'].number_format='0.00'
   chart=LineChart();chart.title='Five-year frequency rates / per 1M hours';chart.add_data(Reference(ws,min_col=6,max_col=8,min_row=1,max_row=len(rows)+1),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=len(rows)+1));chart.width=25;chart.height=12;ws.add_chart(chart,'K2')
 summary.column_dimensions['A'].width=35;summary.column_dimensions['B'].width=70
 for row in summary:
  for cell in row:cell.font=Font(name='Arial',size=12,color='163747')
 wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True);out=io.BytesIO();wb.save(out);return out.getvalue()

def complete_pptx(d):
 from pptx import Presentation
 from pptx.util import Inches,Pt
 from pptx.dml.color import RGBColor
 prs=Presentation(io.BytesIO(summary_pptx(d)));layout=prs.slide_layouts[6]
 def text(s,x,y,w,h,v,size=16,bold=False,color='163747'):
  t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));t.text_frame.word_wrap=True;p=t.text_frame.paragraphs[0];p.text=str(v);p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);return t
 def base(title):
  s=prs.slides.add_slide(layout);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F5F8FA');text(s,.5,.22,15,.35,str(d.get('brand','NAVO'))+' / '+str(d.get('scope','')),12,True,'167563');text(s,.5,.8,15,.75,title,30,True);text(s,.5,8.35,14,.3,('DEMO / ' if d.get('illustrative') else '')+str(d.get('period','')),10,False,'607986');return s
 cover=base(d.get('dashboardTitle','Leadership')+' / executive brief');text(cover,.5,2.4,14,1.5,'One clear view.\nA concise summary of the selected dashboard.',32,True);text(cover,.5,4.7,14,.7,d.get('period','')+' / '+d.get('country',''),20);text(cover,.5,6.2,14,.8,'Illustrative demonstration data' if d.get('illustrative') else 'Accessible reporting records',18,False,'976B13');link=text(cover,.5,7.3,14,.4,'Open live dashboard',18,True,'167563')
 for run in link.text_frame.paragraphs[0].runs:run.hyperlink.address=d.get('sourceUrl','https://navo-hse.vercel.app/')
 slide_id=prs.slides._sldIdLst[-1];prs.slides._sldIdLst.remove(slide_id);prs.slides._sldIdLst.insert(0,slide_id)
 # Slide 2 is the actual full overview captured by the signed-in browser.
 b=image_bytes(d)
 if b:
  s=prs.slides[1]
  for sh in list(s.shapes):sh._element.getparent().remove(sh._element)
  text(s,.35,.12,15,.5,d.get('dashboardTitle','Dashboard')+' / '+str(d.get('period','')),20,True)
  from PIL import Image
  im=Image.open(io.BytesIO(b));w,h=15.3,7.85;r=min(w/im.width,h/im.height);iw,ih=im.width*r,im.height*r;s.shapes.add_picture(io.BytesIO(b),Inches((16-iw)/2),Inches(.8+(h-ih)/2),width=Inches(iw),height=Inches(ih))
 for title,headers,rows in sections(d):
  chunks=[rows[i:i+10] for i in range(0,len(rows),10)] or [[]]
  for idx,part in enumerate(chunks):
   s=base(title+(f' / {idx+1} of {len(chunks)}' if len(chunks)>1 else ''));text(s,.5,1.65,14,.35,f'{len(rows)} total rows / selected scope',12,False,'607986')
   table=s.shapes.add_table(len(part)+1,len(headers),Inches(.5),Inches(2.15),Inches(15),Inches(min(5.8,.44*(len(part)+1)))).table
   for i,row in enumerate([headers]+part):
    for j,v in enumerate(row):
     cell=table.cell(i,j);cell.text=value(v);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('163747' if i==0 else ('FFFFFF' if i%2 else 'EDF3F5'));cell.margin_left=Inches(.07);cell.margin_right=Inches(.07)
     for p in cell.text_frame.paragraphs:p.font.name='Arial';p.font.size=Pt(11 if len(headers)<7 else 10);p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '163747');p.font.bold=i==0
   if title in ['Source record register','Action register']:
    widths=[1,5.5,2.4,2.6,1.5,2] if len(headers)==6 else [1,5,2.5,2,1.5,1.6,1.4]
    for col,w in zip(table.columns,widths):col.width=Inches(w)
 out=io.BytesIO();prs.save(out);return out.getvalue()
