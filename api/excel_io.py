"""Excel template and reader for monthly performance and environmental data.

template_xlsx(d): d = {"month": "YYYY-MM", "projects": [{"id","name"}], "monthly": [[key,label],...], "env": [[key,label],...], "existing": {...}}
parse_xlsx(raw, d): returns {"rows": [...], "errors": [...]} without writing anything.
"""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, TEAL, PALE, GREY = "14325C", "0E9E8E", "F3F6FA", "5F6F84"
SHEETS = (("Monthly data", "monthly"), ("Environmental", "env"))


def template_xlsx(d):
    wb = Workbook()
    info = wb.active; info.title = "Read me"
    lines = ["NAVO data import template", "",
             "1. Fill one row per project and month on the 'Monthly data' and 'Environmental' sheets.",
             "2. Keep the Project ID and Month columns as they are. Month format: YYYY-MM.",
             "3. Leave a cell blank when the figure was not measured. Enter 0 only for a confirmed zero.",
             "4. Numbers only, no units. Negative numbers are rejected.",
             "5. Upload the file in Administration → Performance data → Import from Excel. You will see a preview before anything is saved.",
             "Existing figures for the same project and month are updated; blank cells never erase saved figures."]
    for i, t in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=t); c.font = Font(name="Arial", size=14 if i == 1 else 11, bold=i == 1, color=NAVY if i == 1 else "122033")
    info.column_dimensions["A"].width = 120
    for title, key in SHEETS:
        ws = wb.create_sheet(title)
        fields = d.get(key) or []
        head = ["Project ID", "Project", "Month"] + [f[1] for f in fields]
        for i, h in enumerate(head, 1):
            c = ws.cell(row=1, column=i, value=h); c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY if i <= 3 else TEAL)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            if i > 3:
                c.comment = Comment(f"Field key: {fields[i - 4][0]}", "NAVO")
        ws.row_dimensions[1].height = 48
        ws.cell(row=2, column=1, value="__keys__").font = Font(color="FFFFFF", size=6)
        for i, f in enumerate(fields, 4):
            ws.cell(row=2, column=i, value=f[0]).font = Font(color=GREY, size=8, italic=True)
        ws.row_dimensions[2].hidden = True
        ex = (d.get("existing") or {}).get(key) or {}
        for r, p in enumerate(d.get("projects") or [], 3):
            ws.cell(row=r, column=1, value=p["id"]); ws.cell(row=r, column=2, value=p.get("name") or p["id"]); ws.cell(row=r, column=3, value=d.get("month"))
            vals = ex.get(p["id"]) or {}
            for i, f in enumerate(fields, 4):
                v = vals.get(f[0])
                c = ws.cell(row=r, column=i, value=v if isinstance(v, (int, float)) else None); c.number_format = "#,##0.##"
            for i in range(1, 4):
                ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=PALE)
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, error="Enter a number of 0 or more, or leave blank.")
        ws.add_data_validation(dv)
        if fields:
            dv.add(f"D3:{get_column_letter(3 + len(fields))}2000")
        ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 11
        for i in range(4, 4 + len(fields)):
            ws.column_dimensions[get_column_letter(i)].width = 16
        ws.freeze_panes = "D3"
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def parse_xlsx(raw, d):
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        raise ValueError("The file could not be read. Upload the NAVO template saved as .xlsx.")
    known = {p["id"] for p in d.get("projects") or []}
    rows, errors = [], []
    for title, key in SHEETS:
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        data = list(ws.iter_rows(values_only=True))
        if len(data) < 2:
            continue
        keys = list(data[1]) if data[1] and data[1][0] == "__keys__" else None
        allowed = {f[0] for f in d.get(key) or []}
        if not keys:
            errors.append(f"{title}: the hidden field row is missing. Download a new template.")
            continue
        for n, row in enumerate(data[2:], 3):
            if not row or all(v in (None, "") for v in row):
                continue
            pid = str(row[0] or "").strip(); month = str(row[2] or "").strip()[:7]
            if re.match(r"^\d{4}-\d{2}-\d{2}", str(row[2] or "")):
                month = str(row[2])[:7]
            if hasattr(row[2], "strftime"):
                month = row[2].strftime("%Y-%m")
            if pid not in known:
                errors.append(f"{title} row {n}: project '{pid}' is not one of your projects."); continue
            if not re.match(r"^\d{4}-(0[1-9]|1[0-2])$", month):
                errors.append(f"{title} row {n}: month '{row[2]}' must be YYYY-MM."); continue
            vals, bad = {}, False
            for i in range(3, min(len(row), len(keys))):
                k, v = keys[i], row[i]
                if k not in allowed or v in (None, ""):
                    continue
                try:
                    f = float(str(v).replace(",", ""))
                except ValueError:
                    errors.append(f"{title} row {n}: '{v}' is not a number."); bad = True; break
                if f < 0:
                    errors.append(f"{title} row {n}: negative value {f}."); bad = True; break
                vals[k] = int(f) if f.is_integer() else round(f, 4)
            if bad:
                continue
            if vals:
                rows.append({"kind": key, "project": pid, "month": month, "values": vals, "row": n, "sheet": title})
    return {"rows": rows[:5000], "errors": errors[:200]}
