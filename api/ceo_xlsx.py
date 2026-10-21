"""Read the client's "CEO Bi-monthly HSSE Dashboard" workbook into NAVO fields.

Understands both layouts in the workbook:
  - the calculation sheet:  label | YTD (Sep 24) | 1-2 Week | 3-4 Week | YTD (Oct 24) | 1-2 Week | 3-4 Week
  - the data entry sheet:   label | Reporting timeline | YTD (Sep 24) | 1st Biweekly | 2nd Biweekly
A "YTD (Sep 24)" column is the year-to-date total up to the end of September; the fortnight columns after it
belong to the following month. Month columns (dates) are monthly values. Target / Completed are training targets.
Division tables (CCTV cameras, observations per division, per shift, case status) are kept as tables.
"""
import io
import re
from datetime import date, datetime

MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

# (field key, test on the normalised label). First match wins, so specific tests come first.
RULES = [
    ("manhours", lambda s: "manhours" in s and "total" in s), ("manpower", lambda s: s.startswith("averagemanpower")),
    ("inspFindings", lambda s: s.startswith("ohsinspectionobservation")), ("inspections", lambda s: s.startswith("ohsinspection")),
    ("rtac", lambda s: s == "rtac" or "roadtraffic" in s), ("mtc", lambda s: s == "mtc" or s.startswith("medicaltreatment")),
    ("rwc", lambda s: s == "rwc" or s.startswith("restrictedwork")), ("imp", lambda s: s == "imp" or "memberofthepublic" in s),
    ("daysLost", lambda s: "dayslost" in s), ("rlta", lambda s: s.startswith("rlta") or s.startswith("reportablelosttimeaccident")),
    ("rlti", lambda s: s.startswith("rlti") or s.startswith("reportablelosttimeinjur")),
    ("nwrFatal", lambda s: s.startswith("fatalitycase")), ("fat", lambda s: s.startswith("fatalit")),
    ("odc", lambda s: s == "odc" or s.startswith("occupationaldisease")), ("si", lambda s: s == "si" or s.startswith("seriousinjur")),
    ("ncrMinor", lambda s: s == "ncrminor"), ("ncrMajor", lambda s: s == "ncrmajor"), ("ncrOpen", lambda s: s == "ncropen"), ("ncrClosed", lambda s: s == "ncrclosed"),
    ("contraventions", lambda s: "contravention" in s),
    ("wasteGen", lambda s: s == "wastegenerated"), ("wasteDiverted", lambda s: s == "wasterecycled"), ("wasteDisposed", lambda s: s == "wastedisposed"),
    ("wasteRecovered", lambda s: s == "wasterecovered"), ("water", lambda s: s.startswith("waterconsumption")),
    ("electricity", lambda s: s.startswith("electricityconsumption")), ("fuelL", lambda s: s.startswith("fuelconsumption")),
    ("diesel", lambda s: s.startswith("dieselconsumption")), ("renewable", lambda s: s.startswith("renewableenergy")),
    ("ghgNet", lambda s: "lessoffset" in s), ("offsetGhg", lambda s: s.startswith("emissionsoffset")), ("ghg", lambda s: "ghgemission" in s),
    ("emissionReduction", lambda s: s.startswith("emissionreduction")), ("envCompliance", lambda s: s.startswith("environmentalcompliance")),
    ("hospitalized", lambda s: s.startswith("hospitalized") or s.startswith("hospitalised")), ("unfit", lambda s: s.startswith("medicallyunfit")),
    ("referrals", lambda s: "emergencyreferral" in s), ("screening", lambda s: "healthscreening" in s), ("surveillance", lambda s: "healthsurveillance" in s),
    ("medicalCases", lambda s: s.startswith("summaryofmedicalcases")),
    ("campIdxInt", lambda s: "campauditperformanceindex" in s and "internal" in s), ("campIdxExt", lambda s: "campauditperformanceindex" in s and "external" in s),
    ("campInsp", lambda s: "hsseinspection" in s), ("campObs", lambda s: "hsseobservation" in s), ("campIncidents", lambda s: "incidentsatcamp" in s),
    ("grievances", lambda s: "grievance" in s), ("campViolations", lambda s: s.startswith("campviolation")), ("disciplinary", lambda s: s.startswith("disciplinary")),
    ("absconding", lambda s: "runaway" in s or "abscond" in s), ("campAttendees", lambda s: "inductionattendee" in s),
    ("campInductions", lambda s: "campsafetyinduction" in s), ("campDrills", lambda s: "mockdrill" in s),
    ("ioshWS", lambda s: s.startswith("ioshwork")), ("ioshMS", lambda s: s.startswith("ioshmanag")),
    ("patrols", lambda s: "patrol" in s), ("secInsp", lambda s: s.startswith("securityinspection")), ("secIncidents", lambda s: s.startswith("securityincident")),
    ("escortScrap", lambda s: "scrapyard" in s), ("escortHospital", lambda s: s.startswith("escorttohospital")),
    ("liftPlans", lambda s: "completedliftplan" in s), ("liftsOk", lambda s: "withoutincident" in s), ("liftsInc", lambda s: "liftscompletedwithincident" in s),
    ("liftsUnsafe", lambda s: "unsafeact" in s), ("liftTraining", lambda s: "trainingscompleted" in s),
]
PERCENT = {"campIdxInt", "campIdxExt", "emissionReduction"}
TABLES = {"cctvcameraoperations": "cctvCams", "noofobservationsperdivision": "obsDivision", "noofobservationsfordivision": "obsDivision",
          "noofobservationspershift": "obsShift", "casestatus": "obsStatus"}


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def key_for(label):
    s = norm(label)
    if not s:
        return None
    for k, test in RULES:
        if test(s):
            return k
    return None


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"-?[\d,]*\.?\d+", str(v))
    return float(m.group().replace(",", "")) if m else None


def ym(y, m):
    return f"{y:04d}-{m:02d}"


def shift(month, n):
    y, m = map(int, month.split("-"))
    m += n
    while m > 12:
        y, m = y + 1, m - 12
    while m < 1:
        y, m = y - 1, m + 12
    return ym(y, m)


def head_kind(v):
    """Classify a header cell: ('month', 'YYYY-MM') | ('ytd', 'YYYY-MM') | ('w1',) | ('w2',) | ('target',) | ('done',) | None"""
    if isinstance(v, (datetime, date)):
        return ("month", ym(v.year, v.month))
    s = str(v or "").strip().lower()
    if not s:
        return None
    if "ytd" in s:
        m = re.search(r"([a-z]{3})[a-z]*\W*(\d{2,4})", s.split("ytd", 1)[1])
        if m and m.group(1) in MONTHS:
            y = int(m.group(2))
            return ("ytd", ym(y + 2000 if y < 100 else y, MONTHS[m.group(1)]))
        return ("ytd", None)
    m = re.search(r"\bto\s+([a-z]{3})[a-z]*\W*(\d{2,4})$", s)
    if m and m.group(1) in MONTHS:  # "July to Oct 24" is a running total up to October
        y = int(m.group(2))
        return ("ytd", ym(y + 2000 if y < 100 else y, MONTHS[m.group(1)]))
    if re.search(r"(1\s*-\s*2\s*week|1st\s*bi)", s):
        return ("w1",)
    if re.search(r"(3\s*-\s*4\s*week|2nd\s*bi)", s):
        return ("w2",)
    if s == "target":
        return ("target",)
    if s in ("completed", "actual"):
        return ("done",)
    m = re.match(r"([a-z]{3})[a-z]*\W*(\d{2,4})$", s)
    if m and m.group(1) in MONTHS:
        y = int(m.group(2))
        return ("month", ym(y + 2000 if y < 100 else y, MONTHS[m.group(1)]))
    return None


def parse(raw):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("This is not an Excel (.xlsx) file.")
    periods, monthly, targets, tables, labels = {}, {}, {}, {}, set()

    def put(month, half, k, v):
        p = periods.setdefault(month, {"opening": {}, "h1": {}, "h2": {}})
        p["h%d" % half if half else "opening"].setdefault(k, v)  # the first sheet that reports a figure wins

    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        cols, table, tcols = None, None, None
        for row in rows:
            cells = row + [None] * 3
            lab_i = next((i for i, c in enumerate(cells[:4]) if isinstance(c, str) and c.strip()), None)
            if lab_i is None:
                continue
            label = cells[lab_i]
            heads = [(i, head_kind(c)) for i, c in enumerate(cells) if i > lab_i]
            heads = [(i, h) for i, h in heads if h]
            t = next((v for k2, v in TABLES.items() if norm(label).startswith(k2)), None)
            if t:  # a division table starts; its header may be on this row or the next
                table, tcols = t, [(i, h[1]) for i, h in heads if h[0] == "month"] or None
                tables.setdefault(table, {})
                if t == "obsShift" or t == "obsStatus":
                    cols = None
                continue
            if table and norm(label) == "division":
                tcols = [(i, h[1]) for i, h in heads if h[0] == "month"] or [(lab_i + 1, "count")]
                continue
            if (len(heads) >= 2 or "reportingtimeline" in "".join(norm(c) for c in cells)) and not key_for(label):  # a section header row such as "SAFETY | YTD (Sep 24) | 1-2 Week ..."
                table = None
                cols, cur = [], None
                for i, h in heads:
                    if h[0] == "ytd":
                        cur = h[1]
                        cols.append((i, "opening", cur))
                    elif h[0] in ("w1", "w2"):
                        cols.append((i, h[0], shift(cur, 1) if cur else None))
                    else:
                        cols.append((i, h[0], h[1] if len(h) > 1 else None))
                continue
            if table:
                name = str(label).strip()
                tab = tables[table].setdefault(name, {})
                for i, m in (tcols or [(lab_i + 1, "count")]):
                    v = num(cells[i]) if i < len(cells) else None
                    if v is not None:
                        tab.setdefault(m, v)
                continue
            k = key_for(label)
            if not k or not cols:
                continue
            labels.add(k)
            for i, kind, month in cols:
                v = num(cells[i]) if i < len(cells) else None
                if v is None:
                    continue
                if k in PERCENT and v <= 1.5:
                    v = round(v * 100, 2)
                if kind == "opening" and month:
                    put(month, 0, k, v)
                elif kind in ("w1", "w2") and month:
                    put(month, 1 if kind == "w1" else 2, k, v)
                elif kind == "month" and month:
                    monthly.setdefault(month, {}).setdefault(k, v)
                elif kind == "target":
                    targets.setdefault(k + "Target", v)
                elif kind == "done":
                    targets.setdefault(k, v)
    charts(raw, monthly, targets, tables, max(periods) if periods else None)
    if not labels and not tables:
        raise ValueError("No CEO dashboard figures were found. Use the CEO bi-monthly HSSE workbook.")
    # a fortnight block whose cells are all empty is not a reported period
    periods = {m: p for m, p in periods.items() if any(p.values())}
    tables = {k: v for k, v in tables.items() if v}
    return {"periods": periods, "monthly": monthly, "targets": targets, "tables": tables, "fields": sorted(labels)}


def _chart_series(xml):
    out = []
    for ser in re.findall(r"<c:ser>(.*?)</c:ser>", xml, re.S):
        tx = re.search(r"<c:tx>.*?<c:v>([^<]*)</c:v>", ser, re.S)
        cat, val = re.search(r"<c:cat>(.*?)</c:cat>", ser, re.S), re.search(r"<c:val>(.*?)</c:val>", ser, re.S)
        vs = lambda b: re.findall(r"<c:v>([^<]*)</c:v>", b.group(1)) if b else []
        out.append(((tx.group(1) if tx else "").strip(), [c.strip() for c in vs(cat)], [num(v) for v in vs(val)]))
    return out


def charts(raw, monthly, targets, tables, last):
    """Charts fed from other workbooks keep their figures in the chart itself; read those caches."""
    import zipfile
    try:
        z = zipfile.ZipFile(io.BytesIO(raw))
    except Exception:
        return
    year = int(last[:4]) if last else None
    upto = last
    for n in z.namelist():
        if not re.match(r"xl/charts/chart\d+\.xml$", n):
            continue
        xml = z.read(n).decode("utf8", "ignore")
        title = norm(" ".join(re.findall(r"<a:t>([^<]*)</a:t>", xml)[:4]))
        S = _chart_series(xml)
        series_key = next((k for t, k in (("electricityconsumption", "electricity"), ("waterconsumption", "water"), ("fuelconsumption", "fuelL")) if t in title), None)
        if series_key and year and S:
            _, cats, vals = S[0]
            pairs = [(MONTHS.get(c[:3].lower()), v) for c, v in zip(cats, vals)]
            lastn = max((m for m, v in pairs if m and v), default=0)
            for m, v in pairs:
                if not m or v is None or m > lastn:
                    continue  # months not reached yet are left blank, not zero
                mon = ym(year, m)
                monthly.setdefault(mon, {}).setdefault(series_key, v)
        elif "environmentaltraining" in title:
            for name, cats, vals in S:
                if "training" in norm(name):
                    for c, v in zip(cats, vals):
                        if v is not None:
                            targets.setdefault("envTrainTarget" if norm(c) == "target" else "envTrain", v)
        elif "environmentalaudits" in title:
            for name, cats, vals in S:
                if vals and vals[0] is not None:
                    targets.setdefault("envAuditsTarget" if norm(name) == "target" else "envAudits", vals[0])
        elif "ncrsytd" in title:
            for name, cats, vals in S:
                k = "envNcrClosed" if "closed" in norm(name) else "envNcrOpen" if "open" in norm(name) else "envNcrTotal"
                if vals and vals[0] is not None:
                    targets.setdefault(k, vals[0])
        elif "observtionsytd" in title or "observationsytd" in title:
            for name, cats, vals in S:
                k = "envObsClosed" if "closed" in norm(name) else "envObsOpen" if "open" in norm(name) else "envObsTotal"
                if vals and vals[0] is not None:
                    targets.setdefault(k, vals[0])
        elif "roadtrafficaccident" in title and S:
            _, cats, vals = S[0]
            t = tables.setdefault("rtacEntity", {})
            for c, v in zip(cats, vals):
                if v is not None:
                    t.setdefault(c, {"count": v})
