"""NAVO account management (Vercel Python function). Only a System admin can call it.

POST /api/users  {"op": ...}  with "Authorization: Bearer <signed-in admin token>"
  template                         -> {"file": base64 .xlsx}   the accounts upload template
  parse    {"file": base64 .xlsx}  -> {"rows": [...]}          raw rows read from the template (the page validates them)
  list                             -> {"users": [...]}         every account with last sign-in and status
  create   {"users": [{email, name, password?}]}  -> {"results": [...]}  new or existing accounts; generated passwords returned once
  update   {"id", "email"?, "name"?, "password"?, "disabled"?}
  delete   {"id"}

Access (position, level, projects) is saved by the page through the admin's own database permissions.
Needs NAVO_SUPABASE_SERVICE_KEY in Vercel. The key never leaves the server.
"""
import base64
import io
import json
import logging
import os
import re
import secrets
import urllib.error
import urllib.parse
import urllib.request
from http.server import BaseHTTPRequestHandler

URL = os.environ.get("NAVO_SUPABASE_URL", "https://gvktcleqiidtmgqxbzsy.supabase.co")
PUB = os.environ.get("NAVO_SUPABASE_PUBLISHABLE_KEY", "sb_publishable_9N-3NV87mniwgZB0dw3IVA_bqK3ppNb")
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
HEAD = ["Full name", "Email", "Position", "Level", "Scope (platform / division / cluster name)", "Projects (codes, comma separated)", "Password (optional)"]
ALPHA = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnpqrstuvwxyz23456789"


class ApiError(Exception):
    def __init__(self, status, msg):
        super().__init__(msg)
        self.status = status


def _key():
    k = os.environ.get("NAVO_SUPABASE_SERVICE_KEY") or next((v for n, v in os.environ.items() if n.strip() == "NAVO_SUPABASE_SERVICE_KEY"), None)
    k = (k or "").strip()
    if not k:
        raise ApiError(503, "Account management is not set up yet. Add NAVO_SUPABASE_SERVICE_KEY in Vercel, then redeploy.")
    return k


def _http(path, method="GET", body=None, key=None, extra=None):
    key = key or _key()
    h = {"apikey": key, "Authorization": "Bearer " + key, "Content-Type": "application/json"}
    h.update(extra or {})
    req = urllib.request.Request(URL + path, method=method, headers=h, data=None if body is None else json.dumps(body).encode())
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            raw = r.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as e:
        try:
            msg = json.loads(e.read() or b"{}")
        except Exception:
            msg = {}
        raise ApiError(e.code, msg.get("msg") or msg.get("message") or msg.get("error_description") or "Request failed.")


def admin_uid(token):
    """Return the caller's user id if they are the owner or a System admin, else None."""
    if not token:
        return None
    try:
        user = _http("/auth/v1/user", key=PUB, extra={"Authorization": "Bearer " + token})
    except ApiError:
        return None
    uid = (user or {}).get("id")
    if not uid:
        return None
    q = urllib.parse.quote(uid)
    if _http(f"/rest/v1/owner?user_id=eq.{q}&select=user_id") or _http(f"/rest/v1/members?user_id=eq.{q}&role=eq.admin&select=user_id"):
        return uid
    return None


def gen_password(n=12):
    while True:
        p = "".join(secrets.choice(ALPHA) for _ in range(n))
        if any(c.isdigit() for c in p) and any(c.isupper() for c in p) and any(c.islower() for c in p):
            return p


def check_password(p):
    if len(p) < 8:
        raise ApiError(400, "Passwords need at least 8 characters.")


# ------------------------------------------------------------------ Excel
def template_xlsx(positions=None, levels=None, scopes=None, projects=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14325C")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for col, w in zip("ABCDEFG", (26, 32, 30, 14, 30, 30, 20)):
        ws.column_dimensions[col].width = w
    ws.append(["Ahmed Ali", "ahmed.ali@example.com", "HSE engineer / officer", "Project", "", (projects or ["P-001"])[0], ""])
    ws.freeze_panes = "A2"
    li = wb.create_sheet("Lists")
    cols = [("Positions", positions or []), ("Levels", levels or ["Corporate", "Platform", "Division", "Cluster", "Project"]), ("Scopes", scopes or []), ("Project codes", projects or [])]
    for i, (h, vals) in enumerate(cols, start=1):
        li.cell(1, i, h).font = Font(bold=True)
        li.column_dimensions["ABCD"[i - 1]].width = 34
        for r, v in enumerate(vals, start=2):
            li.cell(r, i, v)
    for col, letter in (("C", "A"), ("D", "B")):
        n = len(cols["AB".index(letter)][1]) + 1
        if n > 1:
            dv = DataValidation(type="list", formula1=f"=Lists!${letter}$2:${letter}${n}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}1000")
    howto = wb.create_sheet("How to")
    for line in ("One row per person. Email is the sign-in name.",
                 "Position sets the rights. Level says how much of the organisation they see.",
                 "Corporate: everything. Platform / Division / Cluster: write the name in Scope. Project: write project codes in Projects.",
                 "Leave Password empty and NAVO creates a strong one. You will see it once after upload.",
                 "People who already have an account keep their password; only their access is updated."):
        howto.append([line])
    howto.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def parse_accounts(raw):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ApiError(400, "This is not an Excel (.xlsx) file.")
    ws = wb["Accounts"] if "Accounts" in wb.sheetnames else wb.worksheets[0]
    rows, head = [], None
    for vals in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in vals]
        if head is None:
            low = [c.lower() for c in cells]
            if any("email" in c for c in low):
                head = low
            continue
        if not any(cells):
            continue
        def get(word):
            for i, h in enumerate(head):
                if h.startswith(word) and i < len(cells):
                    return cells[i]
            return ""
        rows.append({"name": get("full name") or get("name"), "email": get("email").lower(), "position": get("position"), "level": get("level"),
                     "scope": get("scope"), "projects": [p.strip() for p in re.split(r"[,;\n]", get("projects")) if p.strip()], "password": get("password")})
        if len(rows) > 1000:
            raise ApiError(400, "Upload up to 1000 people at a time.")
    if head is None:
        raise ApiError(400, "No header row with an Email column was found. Use the NAVO template.")
    return rows


# ------------------------------------------------------------------ accounts
def list_users():
    out, page = [], 1
    while True:
        d = _http(f"/auth/v1/admin/users?page={page}&per_page=200") or {}
        us = d.get("users", d if isinstance(d, list) else [])
        for u in us:
            out.append({"id": u["id"], "email": u.get("email"), "name": (u.get("user_metadata") or {}).get("name") or "",
                        "lastSignIn": u.get("last_sign_in_at"), "created": u.get("created_at"),
                        "disabled": bool(u.get("banned_until")) and u.get("banned_until") > "2000"})
        if len(us) < 200 or page > 50:
            return out
        page += 1


def _people(uid, name, email):
    body = {"user_id": uid, "email": email}
    if name:
        body["name"] = name
    _http("/rest/v1/people?on_conflict=user_id", "POST", [body], extra={"Prefer": "resolution=merge-duplicates,return=minimal"})


def create_users(users):
    existing = {(u["email"] or "").lower(): u["id"] for u in list_users()}
    results = []
    for u in users[:1000]:
        email = (u.get("email") or "").strip().lower()
        name = (u.get("name") or "").strip()
        if not EMAIL.match(email):
            results.append({"email": email, "status": "error", "message": "Email is not valid."})
            continue
        try:
            if email in existing:
                uid = existing[email]
                _people(uid, name, email)
                results.append({"email": email, "id": uid, "status": "existing"})
                continue
            pw = (u.get("password") or "").strip()
            generated = not pw
            pw = pw or gen_password()
            check_password(pw)
            r = _http("/auth/v1/admin/users", "POST", {"email": email, "password": pw, "email_confirm": True, "user_metadata": {"name": name}})
            uid = r["id"]
            existing[email] = uid
            _people(uid, name, email)
            res = {"email": email, "id": uid, "status": "created"}
            if generated:
                res["password"] = pw
            results.append(res)
        except ApiError as e:
            results.append({"email": email, "status": "error", "message": str(e)})
    return results


def update_user(d, caller):
    uid = d.get("id")
    if not uid:
        raise ApiError(400, "No account selected.")
    body, out = {}, {}
    if d.get("email"):
        e = d["email"].strip().lower()
        if not EMAIL.match(e):
            raise ApiError(400, "Email is not valid.")
        body["email"], body["email_confirm"] = e, True
    if "name" in d:
        body["user_metadata"] = {"name": (d.get("name") or "").strip()}
    if d.get("password") is not None:
        pw = d["password"].strip() or gen_password()
        check_password(pw)
        body["password"] = pw
        if not d["password"].strip():
            out["password"] = pw
    if "disabled" in d:
        if d["disabled"] and uid == caller:
            raise ApiError(400, "You cannot disable your own account.")
        body["ban_duration"] = "876000h" if d["disabled"] else "none"
    if body:
        _http(f"/auth/v1/admin/users/{urllib.parse.quote(uid)}", "PUT", body)
    if "name" in d or d.get("email"):
        patch = {}
        if "name" in d:
            patch["name"] = (d.get("name") or "").strip()
        if d.get("email"):
            patch["email"] = body["email"]
        _http(f"/rest/v1/people?user_id=eq.{urllib.parse.quote(uid)}", "PATCH", patch, extra={"Prefer": "return=minimal"})
    out["ok"] = True
    return out


def delete_user(uid, caller):
    if not uid:
        raise ApiError(400, "No account selected.")
    if uid == caller:
        raise ApiError(400, "You cannot delete your own account.")
    if _http(f"/rest/v1/owner?user_id=eq.{urllib.parse.quote(uid)}&select=user_id"):
        raise ApiError(400, "The owner account cannot be deleted.")
    _http(f"/auth/v1/admin/users/{urllib.parse.quote(uid)}", "DELETE")
    return {"ok": True}


def handle(d, caller):
    op = d.get("op")
    if op == "template":
        return {"file": template_xlsx(d.get("positions"), d.get("levels"), d.get("scopes"), d.get("projects"))}
    if op == "parse":
        if not isinstance(d.get("file"), str):
            raise ApiError(400, "No file received.")
        return {"rows": parse_accounts(base64.b64decode(d["file"].split(",")[-1]))}
    if op == "list":
        return {"users": list_users()}
    if op == "create":
        return {"results": create_users(d.get("users") or [])}
    if op == "update":
        return update_user(d, caller)
    if op == "delete":
        return delete_user(d.get("id"), caller)
    if op == "requests":
        rows = _http("/rest/v1/access_requests?select=*&order=created_at.desc&limit=200") or []
        return {"requests": rows}
    if op == "request_decide":
        rid, status = str(d.get("id") or ""), d.get("status")
        if status not in ("Approved", "Rejected") or not re.match(r"^[0-9a-f-]{36}$", rid):
            raise ApiError(400, "Bad request.")
        _http("/rest/v1/access_requests?id=eq." + rid, "PATCH", {"status": status, "decided_by": caller, "decided_at": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat(), "note": str(d.get("note") or "")[:300]}, extra={"Prefer": "return=minimal"})
        return {"ok": True}
    if op == "qr_tokens":
        import hashlib, hmac
        return {"tokens": {str(p): hmac.new(_key().encode(), ("navo-qr:" + str(p)).encode(), hashlib.sha256).hexdigest()[:20] for p in (d.get("projects") or [])[:500]}}
    raise ApiError(400, "Unknown request.")


class handler(BaseHTTPRequestHandler):
    def _send(self, status, payload):
        b = json.dumps(payload).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def do_POST(self):
        try:
            n = int(self.headers.get("content-length") or 0)
            if n <= 0 or n > 8_000_000:
                return self._send(413, {"error": "The request is too large."})
            d = json.loads(self.rfile.read(n))
            _key()
            auth = self.headers.get("authorization", "")
            caller = admin_uid(auth[7:] if auth.lower().startswith("bearer ") else "")
            if not caller:
                return self._send(401, {"error": "Only a System admin can manage accounts."})
            self._send(200, handle(d, caller))
        except ApiError as e:
            self._send(e.status if 400 <= e.status < 600 else 400, {"error": str(e)})
        except Exception:
            logging.exception("NAVO users failed")
            self._send(500, {"error": "The request failed."})
